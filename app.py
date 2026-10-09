import csv
import io
import json
import os
import traceback
from datetime import datetime

from flask import (Flask, render_template, request, Response,
                   redirect, url_for, flash, jsonify, send_file)
from flask_sqlalchemy import SQLAlchemy

from tasks import TASKS, TOTAL_TASKS


# ---------------------------------------------------------------------------
# App setup
# ---------------------------------------------------------------------------
app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "dev-secret-change-me")

db_url = os.environ.get("DATABASE_URL", "sqlite:///database.db")
if db_url.startswith("postgres://"):
    db_url = db_url.replace("postgres://", "postgresql://", 1)
app.config["SQLALCHEMY_DATABASE_URI"] = db_url
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

db = SQLAlchemy(app)


# ---------------------------------------------------------------------------
# Models
# ---------------------------------------------------------------------------
class Participant(db.Model):
    id             = db.Column(db.Integer, primary_key=True)
    participant_id = db.Column(db.String(50), nullable=False)
    age            = db.Column(db.Integer, nullable=False)
    device         = db.Column(db.String(50), nullable=False)
    vision_group   = db.Column(db.String(50), nullable=False)
    created_at     = db.Column(db.DateTime, default=datetime.utcnow)

    responses = db.relationship("TaskResponse", backref="participant",
                                cascade="all, delete-orphan", lazy=True)


class TaskResponse(db.Model):
    id             = db.Column(db.Integer, primary_key=True)
    participant_id = db.Column(db.Integer, db.ForeignKey("participant.id"),
                               nullable=False)
    task_number    = db.Column(db.Integer, nullable=False)
    task_name      = db.Column(db.String(100), nullable=False)
    response       = db.Column(db.String(2000))
    data           = db.Column(db.Text)
    correct        = db.Column(db.Boolean)
    skipped        = db.Column(db.Boolean, default=False)
    duration_ms    = db.Column(db.Integer)
    completed_at   = db.Column(db.DateTime, default=datetime.utcnow)


# ---------------------------------------------------------------------------
# Boot: create tables + auto-migrate
# ---------------------------------------------------------------------------
with app.app_context():
    db.create_all()
    try:
        from sqlalchemy import text, inspect
        insp = inspect(db.engine)
        if "task_response" in insp.get_table_names():
            existing = {c["name"] for c in insp.get_columns("task_response")}
            migrations = [
                ("data",        "ALTER TABLE task_response ADD COLUMN data TEXT"),
                ("skipped",     "ALTER TABLE task_response ADD COLUMN skipped BOOLEAN DEFAULT FALSE"),
                ("duration_ms", "ALTER TABLE task_response ADD COLUMN duration_ms INTEGER"),
                ("correct",     "ALTER TABLE task_response ADD COLUMN correct BOOLEAN"),
            ]
            for col, sql in migrations:
                if col not in existing:
                    try:
                        db.session.execute(text(sql))
                        db.session.commit()
                        print(f"[migrate] added column: {col}")
                    except Exception as e:
                        db.session.rollback()
                        print(f"[migrate] {col} skipped: {e}")
    except Exception as e:
        print(f"[migrate] failed: {e}")


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------
@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "POST":
        pid = (request.form.get("participantId") or "").strip()
        age = (request.form.get("age") or "").strip()
        dev = request.form.get("device", "Tablet")
        vg  = request.form.get("visionGroup", "Low vision")

        if not pid or not age:
            flash("Please fill in the Participant ID and Age.", "error")
            return redirect(url_for("index"))
        try:
            age_int = int(age)
        except ValueError:
            flash("Age must be a number.", "error")
            return redirect(url_for("index"))

        p = Participant(participant_id=pid, age=age_int, device=dev, vision_group=vg)
        db.session.add(p)
        db.session.commit()
        return redirect(url_for("assessment_task", participant_id=p.id, task_num=1))

    return render_template("index.html", total_tasks=TOTAL_TASKS)


@app.route("/assessment/<int:participant_id>/task/<int:task_num>")
def assessment_task(participant_id, task_num):
    participant = Participant.query.get_or_404(participant_id)
    if task_num > TOTAL_TASKS:
        return redirect(url_for("assessment_complete", participant_id=participant.id))
    if task_num < 1:
        return redirect(url_for("assessment_task",
                                participant_id=participant.id, task_num=1))

    task     = TASKS[task_num - 1]
    progress = int((task_num - 1) / TOTAL_TASKS * 100)
    return render_template("task.html",
                           participant=participant, task=task,
                           task_num=task_num, total=TOTAL_TASKS,
                           progress=progress)


@app.route("/api/submit_task", methods=["POST"])
def api_submit_task():
    try:
        payload = request.get_json(silent=True) or {}
        participant_id = payload.get("participant_id")
        task_number    = payload.get("task_number")

        if not participant_id or not task_number:
            return jsonify({"ok": False, "error": "missing participant_id or task_number"}), 400

        participant = Participant.query.get(participant_id)
        if not participant:
            return jsonify({"ok": False,
                            "error": f"participant {participant_id} not found"}), 404

        task = next((t for t in TASKS if t["n"] == task_number), None)
        if not task:
            return jsonify({"ok": False, "error": f"unknown task {task_number}"}), 400

        skipped      = bool(payload.get("skipped"))
        response_val = payload.get("response", "")
        rich_data    = payload.get("data", {})
        duration_ms  = payload.get("duration_ms")

        correct = None
        if not skipped:
            cfg = task.get("config", {})
            t   = task["type"]
            if t == "ishihara":
                correct = str(response_val).strip() == cfg.get("number", "")
            elif t in ("reading", "sign", "label"):
                correct = str(response_val).strip().lower() == cfg.get("text", "").lower()
            elif t == "obstacle":
                correct = (str(response_val).strip().lower() == "yes") == bool(cfg.get("has_obstacle"))

        rec = TaskResponse(
            participant_id=participant.id,
            task_number=task_number,
            task_name=task["name"],
            response=str(response_val)[:2000] if response_val is not None else "",
            data=json.dumps(rich_data)[:10000] if rich_data else None,
            correct=correct,
            skipped=skipped,
            duration_ms=duration_ms,
        )
        db.session.add(rec)
        db.session.commit()

        next_num = task_number + 1
        next_url = (url_for("assessment_complete", participant_id=participant.id)
                    if next_num > TOTAL_TASKS else
                    url_for("assessment_task",
                            participant_id=participant.id, task_num=next_num))
        return jsonify({"ok": True, "next": next_url})

    except Exception as e:
        db.session.rollback()
        traceback.print_exc()
        return jsonify({"ok": False,
                        "error": f"{type(e).__name__}: {e}"}), 500


@app.route("/assessment/<int:participant_id>/complete")
def assessment_complete(participant_id):
    participant = Participant.query.get_or_404(participant_id)
    responses   = (TaskResponse.query
                   .filter_by(participant_id=participant.id)
                   .order_by(TaskResponse.task_number).all())
    scored = [r for r in responses if r.correct is not None and not r.skipped]
    score  = sum(1 for r in scored if r.correct)
    return render_template("complete.html",
                           participant=participant,
                           responses=responses,
                           score=score,
                           scored_total=len(scored))


# ---------------------------------------------------------------------------
# Spreadsheet / Data viewer
# ---------------------------------------------------------------------------
@app.route("/data")
def data_view():
    participants = Participant.query.order_by(Participant.created_at.desc()).all()

    rows = []
    for p in participants:
        resp_by_num = {r.task_number: r for r in p.responses}
        scored = [r for r in p.responses if r.correct is not None and not r.skipped]
        skipped_count = sum(1 for r in p.responses if r.skipped)
        row = {
            "id":             p.id,
            "participant_id": p.participant_id,
            "age":            p.age,
            "device":         p.device,
            "vision_group":   p.vision_group,
            "created_at":     p.created_at.strftime("%Y-%m-%d %H:%M"),
            "tasks_done":     len(p.responses),
            "skipped":        skipped_count,
            "score":          sum(1 for r in scored if r.correct),
            "scored_total":   len(scored),
            "tasks":          {},
        }
        for t in TASKS:
            r = resp_by_num.get(t["n"])
            if r is None:
                row["tasks"][t["n"]] = {"response": "—", "correct": None, "skipped": False}
            else:
                row["tasks"][t["n"]] = {
                    "response": r.response or "",
                    "correct":  r.correct,
                    "skipped":  bool(r.skipped),
                }
        rows.append(row)

    return render_template("data.html",
                           rows=rows,
                           tasks=TASKS,
                           total_participants=len(participants))


@app.route("/data/participant/<int:participant_id>")
def data_participant(participant_id):
    participant = Participant.query.get_or_404(participant_id)
    responses = (TaskResponse.query
                 .filter_by(participant_id=participant.id)
                 .order_by(TaskResponse.task_number).all())
    return render_template("participant.html",
                           participant=participant,
                           responses=responses)


@app.route("/api/delete_participant/<int:participant_id>", methods=["POST"])
def api_delete_participant(participant_id):
    p = Participant.query.get_or_404(participant_id)
    db.session.delete(p)
    db.session.commit()
    return jsonify({"ok": True})


# ---------------------------------------------------------------------------
# CSV helpers
# ---------------------------------------------------------------------------
def _write_long_rows(writer, participants):
    writer.writerow(["Participant ID", "Age", "Device", "Vision Group",
                     "Task #", "Task Name", "Response", "Correct", "Skipped",
                     "Duration (ms)", "Details", "Timestamp"])
    for p in participants:
        for r in sorted(p.responses, key=lambda x: x.task_number):
            writer.writerow([
                p.participant_id, p.age, p.device, p.vision_group,
                r.task_number, r.task_name, r.response or "",
                "" if r.correct is None else ("1" if r.correct else "0"),
                "1" if r.skipped else "0",
                r.duration_ms or "",
                r.data or "",
                r.completed_at.strftime("%Y-%m-%d %H:%M:%S") if r.completed_at else "",
            ])


def _wide_header():
    header = ["Participant ID", "Age", "Device", "Vision Group", "Date",
              "Tasks Done", "Skipped", "Score", "Scored Total"]
    for t in TASKS:
        header += [f"T{t['n']}_Response", f"T{t['n']}_Correct",
                   f"T{t['n']}_Skipped", f"T{t['n']}_ms"]
    return header


def _wide_row(p):
    by_num = {r.task_number: r for r in p.responses}
    scored = [r for r in p.responses if r.correct is not None and not r.skipped]
    row = [p.participant_id, p.age, p.device, p.vision_group,
           p.created_at.strftime("%Y-%m-%d %H:%M"),
           len(p.responses),
           sum(1 for r in p.responses if r.skipped),
           sum(1 for r in scored if r.correct),
           len(scored)]
    for t in TASKS:
        r = by_num.get(t["n"])
        if r:
            row += [r.response or "",
                    "" if r.correct is None else int(r.correct),
                    int(bool(r.skipped)),
                    r.duration_ms or ""]
        else:
            row += ["", "", "", ""]
    return row


@app.route("/export")
def export_csv():
    output = io.StringIO()
    w = csv.writer(output)
    _write_long_rows(w, Participant.query.order_by(Participant.created_at).all())
    output.seek(0)
    return Response(output.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition":
                             "attachment; filename=all_participants.csv"})


@app.route("/export/participant/<int:participant_id>")
def export_participant(participant_id):
    p = Participant.query.get_or_404(participant_id)
    output = io.StringIO()
    w = csv.writer(output)
    _write_long_rows(w, [p])
    output.seek(0)
    return Response(output.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition":
                             f"attachment; filename=participant_{p.participant_id}.csv"})


@app.route("/export/wide")
def export_wide():
    output = io.StringIO()
    w = csv.writer(output)
    w.writerow(_wide_header())
    for p in Participant.query.order_by(Participant.created_at).all():
        w.writerow(_wide_row(p))
    output.seek(0)
    return Response(output.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition":
                             "attachment; filename=wide_format.csv"})


# ---------------------------------------------------------------------------
# Excel (.xlsx) exports
# ---------------------------------------------------------------------------
def _build_excel(participants, filename):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0A0A0A")
    wrap        = Alignment(wrap_text=True, vertical="top")

    def style_header(ws):
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"

    # ---- Sheet 1: Summary ----
    ws1 = wb.active
    ws1.title = "Summary"
    ws1.append(["Participant ID", "Age", "Device", "Vision Group",
                "Date", "Tasks Done", "Skipped", "Score", "Scored Total"])
    for p in participants:
        ws1.append(_wide_row(p)[:9])
    style_header(ws1)
    for col in ws1.columns:
        max_len = max((len(str(c.value)) if c.value else 0) for c in col)
        ws1.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 10), 40)

    # ---- Sheet 2: Wide Format ----
    ws2 = wb.create_sheet("Wide Format")
    ws2.append(_wide_header())
    for p in participants:
        ws2.append(_wide_row(p))
    style_header(ws2)
    for col in ws2.columns:
        ws2.column_dimensions[col[0].column_letter].width = 14

    # ---- Sheet 3: Long Format ----
    ws3 = wb.create_sheet("Long Format")
    ws3.append(["Participant ID", "Age", "Device", "Vision Group",
                "Task #", "Task Name", "Response", "Correct", "Skipped",
                "Duration (ms)", "Details", "Timestamp"])
    for p in participants:
        for r in sorted(p.responses, key=lambda x: x.task_number):
            ws3.append([
                p.participant_id, p.age, p.device, p.vision_group,
                r.task_number, r.task_name, r.response or "",
                "" if r.correct is None else ("Yes" if r.correct else "No"),
                "Yes" if r.skipped else "No",
                r.duration_ms or "",
                r.data or "",
                r.completed_at.strftime("%Y-%m-%d %H:%M:%S") if r.completed_at else "",
            ])
    style_header(ws3)
    widths = [14, 6, 10, 12, 7, 24, 30, 8, 8, 12, 40, 20]
    for i, w in enumerate(widths, start=1):
        ws3.column_dimensions[ws3.cell(row=1, column=i).column_letter].width = w
    for row in ws3.iter_rows(min_row=2):
        for cell in row:
            if cell.column == 11:
                cell.alignment = wrap

    # ---- Sheet 4+: Per-participant (only if <= 20 participants) ----
    if len(participants) <= 20:
        for p in participants:
            sheet_name = f"P_{p.participant_id}"[:31]
            ws = wb.create_sheet(sheet_name)
            ws.append(["#", "Task", "Response", "Result", "Duration (ms)", "Timestamp"])
            for r in sorted(p.responses, key=lambda x: x.task_number):
                result = ("Skipped" if r.skipped
                          else "—" if r.correct is None
                          else "Correct" if r.correct else "Incorrect")
                ws.append([
                    r.task_number, r.task_name, r.response or "", result,
                    r.duration_ms or "",
                    r.completed_at.strftime("%Y-%m-%d %H:%M:%S") if r.completed_at else "",
                ])
            style_header(ws)
            for i, w in enumerate([5, 24, 30, 12, 14, 20], start=1):
                ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, filename


@app.route("/export/excel")
def export_excel():
    participants = Participant.query.order_by(Participant.created_at).all()
    buf, fname = _build_excel(participants, "low_vision_all.xlsx")
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=fname,
    )


@app.route("/export/excel/participant/<int:participant_id>")
def export_excel_participant(participant_id):
    p = Participant.query.get_or_404(participant_id)
    buf, _ = _build_excel([p], f"participant_{p.participant_id}.xlsx")
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"participant_{p.participant_id}.xlsx",
    )


@app.route("/health")
def health():
    return {"status": "ok", "tasks": TOTAL_TASKS}


if __name__ == "__main__":
    app.run(debug=True)